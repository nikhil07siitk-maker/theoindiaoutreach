"""Merge leads + enrichment and write the outreach workbook, CSVs and a Markdown summary.

Run from the repo root:  python3 research/build.py
"""

import csv
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import check_pages as CP  # noqa: E402
import enrichment as E  # noqa: E402
import leads as L  # noqa: E402
import wave2 as W  # noqa: E402
import wave3 as W3  # noqa: E402
import wave4 as W4  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_XLSX = os.path.join(ROOT, "North_West_East_India_Specialty_Coffee_Leads.xlsx")
CSV_DIR = os.path.join(ROOT, "csv")

REGION_ORDER = [L.WEST, L.CENTRAL, L.NORTH, L.EAST, L.NE]
STATE_ORDER = [
    "Maharashtra", "Gujarat", "Goa", "Madhya Pradesh", "Chhattisgarh", "Rajasthan", "Delhi", "Delhi NCR",
    "Delhi NCR (TBC)", "Haryana", "Uttar Pradesh", "Uttar Pradesh / Punjab", "Uttar Pradesh (TBC)", "Punjab",
    "Chandigarh", "Himachal Pradesh", "Uttarakhand", "Ladakh", "West Bengal", "West Bengal / Nagaland", "Odisha",
    "Bihar", "Jharkhand", "Assam", "Meghalaya", "Sikkim", "Nagaland", "Mizoram", "Manipur", "Tripura", "Jammu & Kashmir", "TBC",
]
STATUS_ORDER = [L.UP_, L.NEW26, L.NEW25, L.EST]
APPEND_FIELDS = ("evidence", "verify")


def merge(base, updates):
    out = []
    for row in base:
        row = dict(row)
        upd = updates.get(row["name"], {})
        for k, v in upd.items():
            if k in APPEND_FIELDS and row.get(k):
                row[k] = f"{row[k]} {v}"
            else:
                row[k] = v
        out.append(row)
    unknown = set(updates) - {r["name"] for r in base}
    if unknown:
        raise SystemExit(f"Enrichment keys with no matching lead: {sorted(unknown)}")
    return out


def order_key(row):
    region = REGION_ORDER.index(row["region"]) if row["region"] in REGION_ORDER else 99
    state = STATE_ORDER.index(row["state"]) if row["state"] in STATE_ORDER else 98
    return region, state, row["city"], row.get("priority", "Z"), row.get("tier", ""), row["name"]


def ig_url(handle):
    h = handle.strip()
    if h.startswith("@") and " " not in h and "/" not in h:
        return f"https://www.instagram.com/{h[1:]}/"
    return ""


def first_contact(row):
    if row.get("email"):
        return "Email"
    if row.get("ig", "").startswith("@"):
        return "Instagram DM"
    if row.get("phone"):
        return "Call / WhatsApp"
    if row.get("linkedin"):
        return "LinkedIn"
    if row.get("web"):
        return "Website contact form"
    return "Walk-in / Google Maps listing"


def default_offer(c):
    if c.get("roasts", "").startswith("Yes"):
        return "Green coffee or a limited guest micro-lot collaboration (they roast in-house)."
    if c["tier"] == "Premium":
        return "Rotating single-estate / micro-lot for the filter bar + a single-origin espresso option."
    return "Premium house espresso blend of Indian estates + one rotating single estate + cold-brew lot."


def default_approach(c):
    if c.get("roasts", "").startswith("Yes"):
        return "Don't pitch roasted supply — offer green lots or a co-branded guest micro-lot."
    if c["status"] == L.UP_:
        return ("Pre-launch: send a 3-estate tasting kit and offer a menu-development session before the menu locks; "
                "90-day launch exclusivity in their micro-market; fixed price for the first 6 months.")
    if c["status"] in (L.NEW26, L.NEW25):
        return ("Just opened: offer a guest-lot month + free barista calibration; side-by-side cupping against their current beans; "
                "ask which roaster/estates they use and what they'd change.")
    if c["type"] == L.MULTI:
        return "Multi-roaster: pitch a guest slot — 1 kg bags, story card (farm, altitude, process, cup score), roast date within 7 days."
    if c["type"] == L.CHAIN:
        return "Multi-outlet: pitch central procurement — fixed roast spec, monthly volume pricing, retail bags for the shelf."
    return "Established: position as a second roaster / rotating single estate; walk-in cupping with the head barista."


cafes = merge(L.CAFES, E.CAFE_UPDATES) + [dict(c) for c in E.NEW_CAFES]
cafes = merge(cafes, W.PROFILE_UPDATES) + [dict(c) for c in W.NEW_CAFES_W2]
cafes = merge(cafes, W3.CAFE_UPDATES_W3) + [dict(c) for c in W3.NEW_CAFES_W3]
cafes = merge(cafes, W4.CAFE_UPDATES_W4) + [dict(c) for c in W4.NEW_CAFES_W4]
missing = set(W4.REMOVE_W4) - {c["name"] for c in cafes}
if missing:
    raise SystemExit(f"REMOVE_W4 names with no matching lead: {sorted(missing)}")
cafes = [c for c in cafes if c["name"] not in W4.REMOVE_W4]
roasters = merge(L.ROASTERS, E.ROASTER_UPDATES) + [dict(r) for r in E.NEW_ROASTERS]
roasters = merge(roasters, W.ROASTER_PROFILE_UPDATES) + [dict(r) for r in W.NEW_ROASTERS_W2]
roasters = merge(roasters, W3.ROASTER_UPDATES_W3) + [dict(r) for r in W3.NEW_ROASTERS_W3]
roasters = merge(roasters, W4.ROASTER_UPDATES_W4) + [dict(r) for r in W4.NEW_ROASTERS_W4]
channels = list(L.CHANNELS) + list(E.NEW_CHANNELS) + list(W3.NEW_CHANNELS_W3) + list(W4.NEW_CHANNELS_W4)

names = [c["name"] for c in cafes]
dupes = {n for n in names if names.count(n) > 1}
if dupes:
    raise SystemExit(f"Duplicate cafe names: {sorted(dupes)}")

for c in cafes:
    assert c["tier"] in ("Premium", "Normal"), c["name"]
    assert c["status"] in STATUS_ORDER, c["name"]
    c["ig_url"] = ig_url(c.get("ig", ""))
    c["contact_route"] = first_contact(c)
    c["roster"] = c.get("roster") or c.get("beans") or "Not published"
    if not c.get("supplier"):
        c["supplier"] = "Own roastery" if c.get("roasts", "").startswith("Yes") else W.NOT_PUB
    c["offer"] = c.get("offer") or default_offer(c)
    c["approach"] = c.get("approach") or default_approach(c)
for r in roasters:
    r["ig_url"] = ig_url(r.get("ig", "").split(" ")[0])

# ---- Web check (research/check_pages.py): auto-extracted page evidence, kept in its own columns
web = {}
if os.path.exists(CP.OUT_JSON):
    with open(CP.OUT_JSON, encoding="utf-8") as f:
        web = {(w["kind"], w["name"]): w for w in json.load(f)}
web_pairs = []
for c in cafes:
    w = web.get(("cafe", c["name"]))
    if not w:
        continue
    c.update({k: v for k, v in CP.summary_fields(w).items() if k in ("web_check", "web_contacts", "web_pages")})
    for role, how in (("own", "Named on the cafe's own pages / listing"), ("mention", "Named in an article about the cafe")):
        for h in w.get(role, {}).get("roasters", []):
            web_pairs.append(dict(name=c["name"], city=c["city"], supplier=h["term"], how=how, snippet=h["snippet"], page_url=h["url"]))
for r in roasters:
    w = web.get(("roaster", r["name"]))
    if not w:
        continue
    r["web_check"] = CP.summary_fields(w)["web_check"]
    r["web_cafes"] = ", ".join(x["cafe"] for x in w.get("cafes_named", []))
    for x in w.get("cafes_named", []):
        city = next((c["city"] for c in cafes if c["name"] == x["cafe"]), "")
        web_pairs.append(dict(name=x["cafe"], city=city, supplier=r["name"], how="Roaster's own page names the cafe",
                              snippet=x["snippet"], page_url=x["url"]))
web_pairs.sort(key=lambda p: (p["name"], p["supplier"]))
for p in web_pairs:
    p["review"] = W4.WEB_PAIR_REVIEW.get((p["name"], p["supplier"]), "Not checked yet")

cafes.sort(key=order_key)
roasters.sort(key=lambda r: order_key(dict(r, priority="", tier=r["tier"])))
for i, c in enumerate(cafes, 1):
    c["id"] = f"C{i:03d}"
for i, r in enumerate(roasters, 1):
    r["id"] = f"R{i:03d}"

hot = [c for c in cafes if c["status"] != L.EST]
hot.sort(key=lambda c: (STATUS_ORDER.index(c["status"]), c["priority"], c["tier"] != "Premium", order_key(c)))

CAFE_COLS = [
    ("ID", "id", 7), ("Priority", "priority", 8), ("Status", "status", 18), ("Tier", "tier", 9), ("Name", "name", 30),
    ("Type", "type", 22), ("Region", "region", 10), ("State", "state", 14), ("City", "city", 16), ("Area / address", "area", 32),
    ("Roasts in-house?", "roasts", 14), ("Current coffee roster", "roster", 50), ("Supplier / roaster", "supplier", 32),
    ("Price signals", "price", 22), ("Gap in their roster", "gap", 40), ("What we can offer", "offer", 45),
    ("Approach strategy", "approach", 55), ("Evidence (why this tier)", "evidence", 60), ("Pitch angle", "angle", 40),
    ("Best first contact", "contact_route", 16),
    ("Instagram", "ig", 24), ("Instagram URL", "ig_url", 34), ("Website", "web", 30), ("Email", "email", 28),
    ("Phone", "phone", 22), ("LinkedIn", "linkedin", 30), ("Listing (Zomato / EazyDiner / Reddit)", "listing", 34),
    ("Key people (public)", "people", 30), ("Source", "src", 45), ("Verify / notes", "verify", 40),
]
ROASTER_COLS = [
    ("ID", "id", 7), ("Tier", "tier", 9), ("Roaster", "name", 32), ("Region", "region", 10), ("State", "state", 16),
    ("City", "city", 22), ("Model / outlets", "model", 50), ("Evidence", "evidence", 50), ("Partnership angle", "angle", 40),
    ("Instagram", "ig", 26), ("Instagram URL", "ig_url", 34), ("Website", "web", 30), ("Email", "email", 26), ("Phone", "phone", 18),
    ("LinkedIn", "linkedin", 40), ("Key people (public)", "people", 34), ("Source", "src", 45), ("Verify / notes", "verify", 30),
]
if web:
    CAFE_COLS += [("Web check: roster evidence (auto — verify)", "web_check", 60), ("Contacts on its own pages", "web_contacts", 30),
                  ("Pages read (web check)", "web_pages", 45)]
    ROASTER_COLS += [("Cafes named on its pages (auto — verify)", "web_cafes", 45), ("Web check (auto — verify)", "web_check", 50)]
WEB_PAIR_COLS = [("Cafe", "name", 30), ("City", "city", 14), ("Roaster named", "supplier", 30), ("How it was found", "how", 30),
                 ("Checked by hand", "review", 36), ("Snippet", "snippet", 70), ("Page", "page_url", 45)]
CHANNEL_COLS = [
    ("Category", "category", 18), ("Name", "name", 34), ("City", "city", 24), ("What it is", "what", 50),
    ("How to use it", "angle", 50), ("Instagram", "ig", 26), ("Website", "web", 40), ("Email", "email", 24),
    ("Phone", "phone", 16), ("LinkedIn", "linkedin", 30), ("Source", "src", 50),
]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="3B2A20")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
LINK_FONT = Font(name=FONT, size=10, color="1F4E9E", underline="single")
THIN = Side(style="thin", color="D9D2C9")
BORDER = Border(bottom=THIN)
A_FILL = PatternFill("solid", fgColor="E3F1DF")
PREM_FILL = PatternFill("solid", fgColor="FBEFD9")
UP_FILL = PatternFill("solid", fgColor="FDE2E1")
LINK_KEYS = {"ig_url", "web", "linkedin", "listing", "src", "page_url"}


def write_table(ws, cols, rows):
    for j, (head, _, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=head)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, row in enumerate(rows, 2):
        for j, (_, key, _) in enumerate(cols, 1):
            val = row.get(key, "") or ""
            cell = ws.cell(row=i, column=j, value=val)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            if key in LINK_KEYS and isinstance(val, str) and val.startswith("http"):
                cell.hyperlink = val
                cell.font = LINK_FONT
    ws.freeze_panes = "F2" if cols is CAFE_COLS else "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}"
    ws.row_dimensions[1].height = 30


def highlight_cafes(ws, n):
    last = get_column_letter(len(CAFE_COLS))
    rng = f"A2:{last}{n + 1}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$C2="Upcoming / Coming soon"'], fill=UP_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$B2="A"'], fill=A_FILL))
    ws.conditional_formatting.add(f"D2:D{n + 1}", FormulaRule(formula=['$D2="Premium"'], fill=PREM_FILL))


wb = Workbook()
readme = wb.active
readme.title = "Read Me"

ws = wb.create_sheet("Hot Leads (New & Upcoming)")
write_table(ws, CAFE_COLS, hot)
highlight_cafes(ws, len(hot))

ws = wb.create_sheet("All Cafes (Premium+Normal)")
write_table(ws, CAFE_COLS, cafes)
highlight_cafes(ws, len(cafes))

PROFILE_COLS = [
    ("ID", "id", 7), ("Priority", "priority", 8), ("Status", "status", 18), ("Tier", "tier", 9), ("Name", "name", 28),
    ("City", "city", 14), ("Roasts in-house?", "roasts", 14), ("Current coffee roster", "roster", 55),
    ("Supplier / roaster", "supplier", 32), ("Price signals", "price", 20), ("Gap in their roster", "gap", 40),
    ("What we can offer", "offer", 45), ("Approach strategy", "approach", 60), ("Best first contact", "contact_route", 16),
    ("Instagram", "ig", 22), ("Key people (public)", "people", 28), ("Source", "src", 40),
]
if web:
    PROFILE_COLS += [("Web check: roster evidence (auto — verify)", "web_check", 60), ("Contacts on its own pages", "web_contacts", 30)]
profiled =[c for c in hot + [c for c in cafes if c["status"] == L.EST] if c["priority"] in ("A", "B")]
ws = wb.create_sheet("Lead Profiles & Strategy")
write_table(ws, PROFILE_COLS, profiled)
ws.freeze_panes = "F2"
ws.conditional_formatting.add(f"A2:{get_column_letter(len(PROFILE_COLS))}{len(profiled) + 1}",
                              FormulaRule(formula=['$B2="A"'], fill=A_FILL))

SUPPLY_COLS = [("Supplier / roaster", "supplier", 40), ("Cafe", "name", 32), ("City", "city", 16), ("Status", "status", 18),
               ("Current coffee roster", "roster", 60), ("Source", "src", 50)]
supply = sorted((c for c in cafes if c["supplier"] != W.NOT_PUB), key=lambda c: (c["supplier"], c["name"]))
ws = wb.create_sheet("Supply Map")
write_table(ws, SUPPLY_COLS, supply)
ws.freeze_panes = "B2"

WEB_SHEET = "Web Check — Supplier Leads"
if web:
    ws = wb.create_sheet(WEB_SHEET)
    write_table(ws, WEB_PAIR_COLS, web_pairs)
    ws.freeze_panes = "B2"

ws = wb.create_sheet("Roasters")
write_table(ws, ROASTER_COLS, roasters)
ws.conditional_formatting.add(f"B2:B{len(roasters) + 1}", FormulaRule(formula=['$B2="Premium"'], fill=PREM_FILL))

ws = wb.create_sheet("Channel Partners")
write_table(ws, CHANNEL_COLS, channels)

ws = wb.create_sheet("Excluded (Regular tier)")
write_table(ws, [("Operator", "name", 40), ("Why excluded", "why", 90)], [dict(name=n, why=w) for n, w in list(L.EXCLUDED) + list(W4.REMOVE_W4.items())])

# ---- Region summary with live formulas
ws = wb.create_sheet("Summary by State")
AC = "'All Cafes (Premium+Normal)'"
heads = ["State", "Cafes (total)", "Premium", "Normal", "Priority A", "New / Upcoming", "Roasters"]
for j, h in enumerate(heads, 1):
    cell = ws.cell(row=1, column=j, value=h)
    cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    ws.column_dimensions[get_column_letter(j)].width = 22 if j == 1 else 15
all_states = {c["state"] for c in cafes} | {r["state"] for r in roasters}
present = [s for s in STATE_ORDER if s in all_states]
present += sorted(all_states - set(present))
n_c, n_r = len(cafes) + 1, len(roasters) + 1


def rng(cols, key, sheet, n):
    letter = get_column_letter([k for _, k, _ in cols].index(key) + 1)
    return f"{sheet}!${letter}$2:${letter}${n}"


ST, TI, PR, STA = (rng(CAFE_COLS, k, AC, n_c) for k in ("state", "tier", "priority", "status"))
for i, s in enumerate(present, 2):
    ws.cell(row=i, column=1, value=s)
    ws.cell(row=i, column=2, value=f"=COUNTIF({ST},A{i})")
    ws.cell(row=i, column=3, value=f'=COUNTIFS({ST},A{i},{TI},"Premium")')
    ws.cell(row=i, column=4, value=f'=COUNTIFS({ST},A{i},{TI},"Normal")')
    ws.cell(row=i, column=5, value=f'=COUNTIFS({ST},A{i},{PR},"A")')
    ws.cell(row=i, column=6, value=f'=COUNTIF({ST},A{i})-COUNTIFS({ST},A{i},{STA},"Established")')
    ws.cell(row=i, column=7, value=f"=COUNTIF({rng(ROASTER_COLS, 'state', 'Roasters', n_r)},A{i})")
tot = len(present) + 2
ws.cell(row=tot, column=1, value="Total")
for j in range(2, 8):
    col = get_column_letter(j)
    ws.cell(row=tot, column=j, value=f"=SUM({col}2:{col}{tot - 1})")
for row in ws.iter_rows(min_row=2, max_row=tot):
    for cell in row:
        cell.font = Font(name=FONT, size=10, bold=cell.row == tot)
ws.freeze_panes = "B2"

# ---- City coverage matrix (how deep each city was searched; lead counts are live formulas)
ws = wb.create_sheet("Coverage by City")
cov_heads = [("State", 18), ("City / cluster", 34), ("Search depth", 12), ("Leads (cafes)", 12), ("Premium", 10),
             ("How it was searched", 70), ("Note", 40)]
for j, (h, w) in enumerate(cov_heads, 1):
    cell = ws.cell(row=1, column=j, value=h)
    cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    ws.column_dimensions[get_column_letter(j)].width = w
CITY = rng(CAFE_COLS, "city", AC, n_c)


def city_count(tokens, extra=""):
    parts = []
    for t in tokens:
        crit = f'"{t[1:]}"' if t.startswith("=") else f'"*{t}*"'
        parts.append(f"COUNTIFS({CITY},{crit}{extra})")
    return "=" + "+".join(parts)


for i, (state, city, depth, how, note) in enumerate(W3.COVERAGE, 2):
    tokens = W3.COVERAGE_TOKENS.get(city) or [t.strip() for t in city.split("/")]
    vals = [state, city, depth, city_count(tokens), city_count(tokens, f',{TI},"Premium"'), how, note]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=j, value=v)
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "C2"

# ---- Platform coverage
ws = wb.create_sheet("Platform Coverage")
PLAT = [
    ("Platform", "What was searched", "What it yielded", "Limits"),
    ("Instagram", "Per-city sweeps ('specialty coffee <city>', 'coming soon') and one search per lead name.",
     "Most handles in this file; 'coming soon' / relaunch posts (Coffee Bond, Brewtown, Cafe Serenity, Craft Coffee Siliguri, Lento Pune); multi-roaster bios (Pilgrim, Coffee Therapy Artisan, Lucky's, Daily Daily, Specialty Coffee House); bio phone numbers.",
     "Search index lags new accounts; DMs are the main contact route for small cafes."),
    ("LinkedIn", "Company pages for roasters/chains; founder 'launching our cafe' posts; cafe hiring posts.",
     "Company pages (Blue Tokai, Subko, KC Roasters, Half Light, Toffee, KaCoBa, 93 Degrees); founders; hiring signals (The Coffee Project Nagpur, Blue Tokai Punjab/Haryana).",
     "Small cafes rarely have pages; profiles of individuals listed only where they are public founders/brand pages."),
    ("News / trade press", "Restaurant India, Indian Retailer, Franchise India, World Coffee Portal, Hotelier India, BW Hotelier, ANI/ThePrint.",
     "Dated openings (Oro, Jacob's Brew House, Momoka, Pour Over Santushti, Gram Street, Cocoa & Katli, Stand By Juhu, Coffee Capital).",
     "Press-release driven: favours funded brands."),
    ("Monthly openings roundups (wave 2)", "Foodies of India (Mumbai, Oct 2025–Oct 2026 monthly), Knocksense (Lucknow / Indore / Ahmedabad), The Lab Mag + Curly Tales + The Nod (Delhi NCR monthly), t2 (Kolkata cafe alerts).",
     "~60 more 2025–26 openings (JOKER, Rola, Brevé, KĀPIBAR, Grahm, Tsubaki, MUVN, Belsta, Javaphile, Lizzu, Off Mrkt, Kokoy, One PourCent, Batter Bar, Cháfé Braz…).",
     "Small new cafes rarely publish their roaster — profiled as 'Not published' with an ask-on-contact strategy."),
    ("City food media", "The Nod, Outlook Traveller (OT Pours), Curly Tales, The Lab Mag, LBB, Time Out, Esquire, Free Press Journal, t2 (Kolkata), Elle Decor.",
     "Monthly 'new openings' lists: The Drop, Therefore, In Good Co., ATE Omakase, Odd Culture, Traut, Flameback, Lento Pune.",
     "Covers metros far better than Tier-2 cities."),
    ("District / Zomato / EazyDiner / Tripadvisor", "City specialty guides, listings for addresses and ratings.",
     "Addresses, ratings, price-for-two (Toise, Traut, Pull and Pour, Kulture, Pilgrim).",
     "Phone numbers mostly behind app; not scraped."),
    ("Reddit", "r/PuneCoffee and city coffee threads.", "Pre-launch founder post (Viman Nagar, Pune).", "Few city subs are active."),
    ("Job boards", "Glassdoor, Apna, Foundit, Weekday barista listings.",
     "Hiring = opening soon signal (Kerb Coffee Co., Basque by breve, Bhatia Speciality Coffee LLP, 729 Grams).", "Employer names sometimes legal entities."),
    ("Directories", "Indian Coffee Beans, Indian Coffee Guide, Speciality Coffee Online, BrewAtlas, Wanderlog.",
     "Roaster lists by city; wholesale terms of competitors (BROOT, Rossette).", "Some directory entries could not be corroborated (flagged)."),
    ("Email", "Searched '<name> email contact' for priority leads.",
     "Few cafes publish email (Artsy, PICO, Roastery Coffee House). Instagram DM / phone is the realistic first touch.",
     "No email has been guessed or constructed — blank means not published."),
]
for i, row in enumerate(PLAT, 1):
    for j, v in enumerate(row, 1):
        cell = ws.cell(row=i, column=j, value=v)
        cell.font = HEAD_FONT if i == 1 else BODY_FONT
        if i == 1:
            cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for j, w in enumerate([22, 55, 70, 45], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# ---- Read Me
lines = [
    ("Premium Indian Specialty Coffee — Outreach Leads (North, West, Central, East & Northeast India)", True),
    ("Research date: 27 Oct 2026. South India deliberately excluded (Goa is included as West coast).", False),
    ("", False),
    ("TIERS", True),
    ("Premium — origin-driven specialty: single-origin / single-estate or micro-lot menus, manual brew bar (V60, Chemex, AeroPress, siphon), multi-roaster or guest-roast programmes, or imports international single origins. Core market for premium beans.", False),
    ("Normal — better than regular: specialty-leaning, espresso-led cafes with a quality house blend, maybe one filter option; food/dessert-led concepts with a deliberate coffee programme. Market for a premium house blend + one rotating estate.", False),
    ("Regular — commodity / chain cafes. Not targeted; listed only in 'Excluded (Regular tier)' for reference.", False),
    ("", False),
    ("PRIORITY", True),
    ("A — act this week: upcoming or just-opened cafes still choosing suppliers, multi-roaster cafes, and cafes known to buy from outside roasters.", False),
    ("B — strong fit: established Premium/Normal independents that do not roast in-house.", False),
    ("C — nurture: roaster-owned cafes (collab/guest lot only), chains, or thin evidence.", False),
    ("", False),
    ("STATUS", True),
    ("Upcoming / Coming soon — announced or 'opening soon' on social/press but not open (the pre-launch hype window).", False),
    ("New (opened 2026) / New (late 2025) — opened recently; supplier relationships still forming.", False),
    ("", False),
    ("HOW TO USE", True),
    ("1. Start with 'Hot Leads (New & Upcoming)' — sorted Upcoming → New, then priority. Row colours: pink = upcoming, green = priority A, cream tier cell = Premium.", False),
    ("1b. 'Lead Profiles & Strategy' — per priority A/B cafe: roasts in-house?, CURRENT COFFEE ROSTER (what they pour today), supplier/roaster, price signals, gap, what we can offer, approach strategy.", False),
    ("1c. 'Supply Map' — every cafe whose supplier or own roastery is publicly known. Where a roster/supplier is 'Not published', ask on first contact and check the retail shelf / bag on the bar.", False),
    ("2. 'Best first contact' picks the most direct public channel found (email > Instagram DM > phone > LinkedIn > website > walk-in).", False),
    ("3. Every row has a Source link. 'Verify / notes' flags single-source facts, unconfirmed cities and handles not yet indexed by search.", False),
    ("4. 'Roasters' lists competitors and possible partners (guest-roast slots in multi-roaster cafes, co-roasting, distribution). 'Channel Partners' lists cafe consultants, barista academies, communities and events that reach cafes before launch.", False),
    ("5. 'Summary by State' recounts automatically (COUNTIF formulas) if you add or retag rows.", False),
    ("", False),
    ("CONTACT DATA POLICY", True),
    ("Only publicly listed business contacts were recorded. Nothing was guessed (no constructed emails). Blank = not published / not found yet. Individual names appear only as publicly credited founders or brand leads.", False),
    ("", False),
]
if web:
    lines += [
        ("WEB CHECK (research/check_pages.py)", True),
        ("Columns and the sheet marked 'auto — verify' were extracted automatically from each lead's own site, listings and articles: roasters named, origins/estates, processes, brew methods, prices, opening/hiring signals and published contacts. Each item has a snippet and page in csv/page_findings.csv and research/page_findings.json. Check by eye before quoting it to the cafe.", False),
        ("'Web Check — Supplier Leads' pairs cafes with roasters named on the cafe's pages, in articles about it, or on the roaster's stockist pages. 'Named in an article' is the weakest evidence.", False),
        ("", False),
    ]
lines += [("COUNTS", True)]
for i, (text, bold) in enumerate(lines, 1):
    cell = readme.cell(row=i, column=1, value=text)
    cell.font = Font(name=FONT, size=13 if i == 1 else 10, bold=bold)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
r0 = len(lines) + 1
HL = "'Hot Leads (New & Upcoming)'"
count_rows = [
    ("Cafes — total (Premium + Normal)", f"=COUNTA({rng(CAFE_COLS, 'id', AC, n_c)})"),
    ("Cafes — Premium", f'=COUNTIF({TI},"Premium")'),
    ("Cafes — Normal", f'=COUNTIF({TI},"Normal")'),
    ("Cafes — Priority A", f'=COUNTIF({PR},"A")'),
    ("Hot leads (new & upcoming)", f"=COUNTA({rng(CAFE_COLS, 'id', HL, len(hot) + 1)})"),
    ("Upcoming / coming soon", f'=COUNTIF({STA},"Upcoming / Coming soon")'),
    ("Opened in 2026", f'=COUNTIF({STA},"New (opened 2026)")'),
    ("Multi-roaster cafes", f'=COUNTIF({rng(CAFE_COLS, "type", AC, n_c)},"Multi-roaster cafe")'),
    ("Cafes that roast in-house", f'=COUNTIF({rng(CAFE_COLS, "roasts", AC, n_c)},"Yes*")'),
    ("Supplier / roaster identified (not own roastery)", f"=COUNTA('Supply Map'!$A$2:$A${len(supply) + 1})-COUNTIF('Supply Map'!$A$2:$A${len(supply) + 1},\"Own roastery*\")"),
    ("With an Instagram handle", f'=COUNTIF({rng(CAFE_COLS, "ig", AC, n_c)},"@*")'),
    ("Roasters", f"=COUNTA({rng(ROASTER_COLS, 'id', 'Roasters', n_r)})"),
    ("Channel partners", f"=COUNTA('Channel Partners'!$A$2:$A${len(channels) + 1})"),
]
if web:
    count_rows += [
        ("Cafes with a web check", f'=COUNTIF({rng(CAFE_COLS, "web_pages", AC, n_c)},"?*")'),
        ("Cafe–roaster pairs from the web check", f"=COUNTA('{WEB_SHEET}'!$A$2:$A${max(len(web_pairs), 1) + 1})"),
    ]
for k, (label, formula) in enumerate(count_rows):
    readme.cell(row=r0 + k, column=1, value=label).font = BODY_FONT
    c = readme.cell(row=r0 + k, column=2, value=formula)
    c.font = Font(name=FONT, size=10, bold=True)
readme.column_dimensions["A"].width = 120
readme.column_dimensions["B"].width = 12

wb.save(OUT_XLSX)

# ---- CSVs
os.makedirs(CSV_DIR, exist_ok=True)


def dump(path, cols, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([h for h, _, _ in cols])
        for row in rows:
            w.writerow([row.get(k, "") or "" for _, k, _ in cols])


dump(os.path.join(CSV_DIR, "hot_leads_new_upcoming.csv"), CAFE_COLS, hot)
dump(os.path.join(CSV_DIR, "cafes_premium_normal.csv"), CAFE_COLS, cafes)
dump(os.path.join(CSV_DIR, "roasters.csv"), ROASTER_COLS, roasters)
dump(os.path.join(CSV_DIR, "channel_partners.csv"), CHANNEL_COLS, channels)
dump(os.path.join(CSV_DIR, "lead_profiles_strategy.csv"), PROFILE_COLS, profiled)
dump(os.path.join(CSV_DIR, "supply_map.csv"), SUPPLY_COLS, supply)
if web:
    dump(os.path.join(CSV_DIR, "web_supplier_leads.csv"), WEB_PAIR_COLS, web_pairs)

# ---- Markdown table of hot leads + external-roaster buyers
def md_cell(v):
    return (v or "").replace("|", "/").replace("\n", " ")


def md_contact(c):
    parts = [c.get("ig", "") if c.get("ig", "").startswith("@") else "", c.get("email", ""), c.get("phone", ""), c.get("web", "")]
    return " · ".join(p for p in parts if p) or "—"


buyers = [c for c in cafes if c["status"] == L.EST and c["priority"] == "A"]
with open(os.path.join(ROOT, "HOT_LEADS.md"), "w", encoding="utf-8") as f:
    f.write("# Hot leads — new & upcoming premium cafes\n\n")
    f.write("Generated by `research/build.py` from `research/leads.py`, `enrichment.py`, `wave2.py`, `wave3.py` and `wave4.py`. "
            "Full profiles (gap, offer, approach, sources, verify notes) are in the workbook's 'Lead Profiles & Strategy' tab.\n\n")
    f.write("| ID | Pri | Status | Tier | Cafe | City | Roasts? | Current coffee roster | Supplier | Contact found |\n"
            "|---|---|---|---|---|---|---|---|---|---|\n")
    for c in hot:
        f.write(f"| {c['id']} | {c['priority']} | {c['status']} | {c['tier']} | {md_cell(c['name'])} | {md_cell(c['city'])} | "
                f"{md_cell(c.get('roasts') or 'Unknown')} | {md_cell(c['roster'])} | {md_cell(c['supplier'])} | {md_cell(md_contact(c))} |\n")
    f.write("\n## Established priority-A cafes (multi-roaster or already buying from outside roasters)\n\n")
    f.write("| ID | Cafe | Type | City | Current coffee roster | Supplier | Contact found |\n|---|---|---|---|---|---|---|\n")
    for c in buyers:
        f.write(f"| {c['id']} | {md_cell(c['name'])} | {c['type']} | {md_cell(c['city'])} | {md_cell(c['roster'])} | "
                f"{md_cell(c['supplier'])} | {md_cell(md_contact(c))} |\n")

print(f"cafes={len(cafes)} hot={len(hot)} roasters={len(roasters)} channels={len(channels)}")
print(f"premium={sum(c['tier'] == 'Premium' for c in cafes)} normal={sum(c['tier'] == 'Normal' for c in cafes)} "
      f"A={sum(c['priority'] == 'A' for c in cafes)} ig={sum(c.get('ig', '').startswith('@') for c in cafes)}")
